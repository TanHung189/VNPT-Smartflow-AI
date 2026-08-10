# ================================================================
# Celery AI Tasks — VNPT SmartFlow AI
#
# Các tác vụ AI nặng được xử lý bất đồng bộ trong nền.
# FastAPI nhận request → tạo task_id → trả về ngay lập tức.
# Worker xử lý → lưu kết quả vào Redis → Frontend polling lấy kết quả.
# ================================================================

import logging
import json
import hashlib
import io
from typing import Optional

import docx
import PyPDF2
import sentry_sdk

from app.core.celery_app import celery_app

logger = logging.getLogger(__name__)


# ─────────────────────────────────────────────────────────────────────────────
# TASK: Tạo sơ đồ từ văn bản (Text-to-Flow)
# ─────────────────────────────────────────────────────────────────────────────
@celery_app.task(
    name="app.tasks.ai_tasks.generate_flow_from_text",
    bind=True,
    max_retries=2,
    default_retry_delay=5,
)
def generate_flow_from_text(
    self,
    text: str,
    provider: str = "gemini",
    the_loai: str = "process",
    user_id: Optional[str] = None,
    current_nodes: Optional[list] = None,
    current_edges: Optional[list] = None,
) -> dict:
    """
    Celery task: Gọi AI để chuyển văn bản thành sơ đồ React Flow.
    Chạy trong Celery Worker (tiến trình riêng biệt).
    """
    import asyncio

    logger.info(f"[Celery Task] generate_flow_from_text | provider={provider} | the_loai={the_loai}")

    try:
        # Xây dựng prompt context từ canvas hiện tại (Chat-to-Edit)
        text_with_context = text
        if current_nodes or current_edges:
            context_str = json.dumps(
                {"nodes": current_nodes or [], "edges": current_edges or []},
                ensure_ascii=False,
            )
            text_with_context = (
                f"{text}\n\n"
                f"[Ngữ cảnh Canvas hiện tại - Hãy ADD/UPDATE/DELETE dựa trên sơ đồ này, "
                f"KHÔNG tạo mới hoàn toàn]:\n{context_str}"
            )

        # Celery worker chạy sync, nên dùng asyncio.run để chạy hàm async
        from app.services.ai_service import ai_service
        result = asyncio.run(
            ai_service.generate_smart_flow(
                text_with_context,
                provider=provider,
                the_loai=the_loai,
                user_id=user_id,
            )
        )
        logger.info(
            f"[Celery Task] ✅ Hoàn thành | "
            f"{len(result.get('nodes', []))} nodes, {len(result.get('edges', []))} edges"
        )
        return {"status": "SUCCESS", "data": result}

    except Exception as exc:
        logger.error(f"[Celery Task] ❌ Lỗi: {exc}", exc_info=True)
        sentry_sdk.capture_exception(exc)
        # Tự động retry nếu còn lượt (ví dụ: Gemini rate limit tạm thời)
        try:
            raise self.retry(exc=exc, countdown=10)
        except self.MaxRetriesExceededError:
            return {"status": "FAILURE", "error": str(exc)}


# ─────────────────────────────────────────────────────────────────────────────
# TASK: Trích xuất file và tạo sơ đồ (File-to-Flow)
# ─────────────────────────────────────────────────────────────────────────────
@celery_app.task(
    name="app.tasks.ai_tasks.generate_flow_from_file",
    bind=True,
    max_retries=1,
    default_retry_delay=10,
)
def generate_flow_from_file(
    self,
    file_content_hex: str,        # bytes được encode hex để truyền qua Celery
    filename: str,
    content_type: str,
    provider: str = "gemini",
    the_loai: str = "process",
    user_id: Optional[str] = None,
) -> dict:
    """
    Celery task: Đọc file (TXT/DOCX/PDF/ảnh) và tạo sơ đồ từ nội dung.
    File content được truyền dưới dạng hex string (safe qua JSON serialization).
    """
    import asyncio
    from app.services.ai_service import ai_service

    logger.info(f"[Celery Task] generate_flow_from_file | file={filename} | provider={provider}")

    try:
        content = bytes.fromhex(file_content_hex)
        filename_lower = filename.lower()
        content_hash = hashlib.md5(content).hexdigest()[:12]

        # ── Ảnh → Vision AI ─────────────────────────────────────────
        if any(filename_lower.endswith(ext) for ext in (".png", ".jpg", ".jpeg", ".webp")):
            logger.info("📸 [Celery] Gọi Vision AI...")
            result = asyncio.run(
                ai_service.generate_flow_from_image(content, content_type, the_loai)
            )
            return {"status": "SUCCESS", "data": result}

        # ── Trích xuất text từ file ──────────────────────────────────
        text = ""
        if filename_lower.endswith(".txt"):
            text = content.decode("utf-8", errors="ignore")
        elif filename_lower.endswith(".docx"):
            doc = docx.Document(io.BytesIO(content))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        parts.append(row_text)
            text = "\n".join(parts)
        elif filename_lower.endswith(".pdf"):
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            text = "\n".join(p.extract_text() or "" for p in reader.pages)

        text = text.strip()
        if not text:
            return {"status": "FAILURE", "error": "Không đọc được nội dung từ file."}

        # ── Xử lý chunk nếu file dài ────────────────────────────────
        CHUNK_LIMIT = 4000
        if len(text) <= CHUNK_LIMIT:
            text_with_hash = f"[FILE_HASH:{content_hash}]\n{text}"
            result = asyncio.run(
                ai_service.generate_smart_flow(
                    text_with_hash,
                    provider=provider,
                    the_loai=the_loai,
                    user_id=user_id,
                )
            )
        else:
            import asyncio as aio

            # Chia chunks theo dòng
            lines = text.splitlines()
            chunks: list = []
            current: list = []
            current_len = 0
            for line in lines:
                line_len = len(line) + 1
                if current_len + line_len > CHUNK_LIMIT and current:
                    chunks.append("\n".join(current))
                    current = [line]
                    current_len = line_len
                else:
                    current.append(line)
                    current_len += line_len
            if current:
                chunks.append("\n".join(current))

            logger.info(f"[Celery Task] File dài → {len(chunks)} chunks xử lý tuần tự")

            # Xử lý từng chunk (trong worker Celery không dùng asyncio.gather tốt)
            merged_nodes: list = []
            merged_edges: list = []

            for i, chunk in enumerate(chunks):
                chunk_text = f"[FILE_HASH:{content_hash}_chunk{i}]\n{chunk}"
                chunk_result = asyncio.run(
                    ai_service.generate_smart_flow(
                        chunk_text,
                        provider=provider,
                        the_loai=the_loai,
                        user_id=user_id,
                    )
                )
                chunk_nodes = chunk_result.get("nodes", [])
                chunk_edges = chunk_result.get("edges", [])
                id_map: dict = {}
                for node in chunk_nodes:
                    old_id = str(node.get("id", ""))
                    new_id = f"c{i}_{old_id}"
                    id_map[old_id] = new_id
                    node["id"] = new_id
                    merged_nodes.append(node)
                for edge in chunk_edges:
                    src = str(edge.get("source", ""))
                    tgt = str(edge.get("target", ""))
                    edge["source"] = id_map.get(src, f"c{i}_{src}")
                    edge["target"] = id_map.get(tgt, f"c{i}_{tgt}")
                    edge["id"] = f"c{i}_{edge.get('id', f'e{src}-{tgt}')}"
                    merged_edges.append(edge)

            result = {"nodes": merged_nodes, "edges": merged_edges}
            logger.info(
                f"[Celery Task] ✅ Merged: {len(merged_nodes)} nodes, {len(merged_edges)} edges"
            )

        return {"status": "SUCCESS", "data": result}

    except Exception as exc:
        logger.error(f"[Celery Task] ❌ Lỗi xử lý file: {exc}", exc_info=True)
        sentry_sdk.capture_exception(exc)
        try:
            raise self.retry(exc=exc, countdown=15)
        except self.MaxRetriesExceededError:
            return {"status": "FAILURE", "error": f"Lỗi xử lý file: {str(exc)}"}
